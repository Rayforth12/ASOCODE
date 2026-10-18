#!/usr/bin/env python3
"""Ensayo antes de crear las cuentas de los asociados (se corre una sola vez, antes de entregar).

Uso:
    pip install openpyxl
    python preparar_asociados.py Listado.xlsx

Revisa el Excel (Id, Cedula, Nombre, Ingreso, Planilla) y escribe en la carpeta actual:
    problemas.csv  filas con errores (no se cargan) y avisos, con el motivo

No sube nada a internet. Cuando no haya errores, siga con crear_cuentas.py.
"""
import csv
import re
import sys
import unicodedata
from collections import defaultdict
from datetime import date, datetime

from openpyxl import load_workbook

COLUMNAS = {"id", "cedula", "nombre", "ingreso", "planilla"}
PLANILLAS = {"semanal", "quincenal", "mensual"}
ID_VALIDO = re.compile(r"^[A-Za-z0-9._-]+$")


def norm(v):
    """Minúsculas, sin tildes ni espacios sobrantes (para comparar encabezados)."""
    s = unicodedata.normalize("NFKD", str(v if v is not None else ""))
    return s.encode("ascii", "ignore").decode().strip().lower()


def texto(v):
    return " ".join(str(v).split()) if v is not None else ""


def fecha(v):
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(texto(v), fmt).date().isoformat()
        except ValueError:
            pass
    return ""


def leer(ruta):
    filas = list(load_workbook(ruta, read_only=True, data_only=True).active.iter_rows(values_only=True))
    for i, f in enumerate(filas[:30]):          # los encabezados pueden no estar en la fila 1
        nombres = [norm(c) for c in f]
        if {"id", "cedula", "nombre"} <= set(nombres):
            idx = {n: nombres.index(n) for n in nombres if n in COLUMNAS}
            break
    else:
        sys.exit("No encontré la fila de encabezados (Id, Cedula, Nombre).")

    regs = []
    for num, f in enumerate(filas[i + 1:], start=i + 2):   # num = número de fila en Excel
        if all(texto(c) == "" for c in f):
            continue
        def g(k):
            return f[idx[k]] if k in idx and idx[k] < len(f) else None
        regs.append({
            "fila": num,
            "tarjeta": texto(g("id")),
            "cedula": re.sub(r"[\s-]", "", texto(g("cedula"))),
            "nombre": texto(g("nombre")),
            "ingreso": fecha(g("ingreso")),
            "planilla": norm(g("planilla")),
        })
    return regs


def validar(regs):
    errores, avisos = defaultdict(list), defaultdict(list)

    for campo in ("tarjeta", "cedula"):                     # repetidos: se marcan todas las filas
        vistos = defaultdict(list)
        for r in regs:
            if r[campo]:
                vistos[r[campo]].append(r["fila"])
        for filas in vistos.values():
            if len(filas) > 1:
                for n in filas:
                    errores[n].append(f"{campo} repetido (filas {', '.join(map(str, filas))})")

    for r in regs:
        n = r["fila"]
        for campo in ("tarjeta", "cedula", "nombre"):
            if not r[campo]:
                errores[n].append(f"falta {campo}")
        if r["cedula"] and not r["cedula"].isdigit():
            errores[n].append("la cédula tiene algo que no es un número")
        elif r["cedula"] and len(r["cedula"]) not in (9, 12):
            avisos[n].append(f"la cédula tiene {len(r['cedula'])} dígitos (se esperan 9 o 12)")
        if r["tarjeta"] and not ID_VALIDO.match(r["tarjeta"]):
            errores[n].append("el Id tiene caracteres que no se pueden usar (solo letras, números, punto y guion)")
        if not r["ingreso"]:
            avisos[n].append("sin fecha de ingreso válida")
        if r["planilla"] not in PLANILLAS:
            avisos[n].append("planilla vacía o no reconocida (queda en blanco)")
            r["planilla"] = ""
    return errores, avisos


def main(ruta):
    regs = leer(ruta)
    errores, avisos = validar(regs)
    buenas = [r for r in regs if r["fila"] not in errores]

    with open("problemas.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["fila_excel", "tipo", "motivo"])
        for n in sorted(set(errores) | set(avisos)):
            for m in errores.get(n, []):
                w.writerow([n, "ERROR (no se carga)", m])
            for m in avisos.get(n, []):
                w.writerow([n, "aviso", m])

    print(f"Filas leídas:                    {len(regs)}")
    print(f"Listas para crear:               {len(buenas)}")
    print(f"Con errores (no se cargan):      {len(errores)}")
    print(f"Filas con avisos (sí se cargan): {len([n for n in avisos if n not in errores])}")
    if errores or avisos:
        print("Detalle en problemas.csv")
    for n in sorted(errores):
        print(f"  Fila {n}: " + "; ".join(errores[n]))


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit("Uso: python preparar_asociados.py Listado.xlsx")
    main(sys.argv[1])
